from __future__ import annotations

import csv
import io
import json
import os
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from typing import Any

MAX_TEXT_CHARS=int(os.getenv("JHADINA_EXTRACTED_TEXT_MAX_CHARS","1000000"))
MAX_ZIP_EXPANDED_BYTES=int(os.getenv("JHADINA_EXTRACTED_ZIP_MAX_BYTES",str(64*1024*1024)))
MAX_SHEET_CELLS=int(os.getenv("JHADINA_EXTRACTED_SHEET_MAX_CELLS","50000"))

DOCX="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MEDIA_MIME={"audio/wav","audio/mpeg","audio/mp4","audio/webm","video/mp4","video/webm"}


class ExtractionError(RuntimeError):
    pass


@dataclass(frozen=True)
class Extraction:
    text:str
    metadata:dict[str,Any]


def _bounded(text:str)->str:
    value=text.strip()
    if not value:
        raise ExtractionError("ARTIFACT_EXTRACTION_EMPTY")
    if len(value)>MAX_TEXT_CHARS:
        value=value[:MAX_TEXT_CHARS]
    return value


def _validate_ooxml(payload:bytes)->None:
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            expanded=sum(info.file_size for info in archive.infolist())
            if expanded>MAX_ZIP_EXPANDED_BYTES:
                raise ExtractionError("ARTIFACT_OOXML_EXPANSION_LIMIT")
    except zipfile.BadZipFile as exc:
        raise ExtractionError("ARTIFACT_OOXML_INVALID_ZIP") from exc


def extract_pdf(payload:bytes)->Extraction:
    from pypdf import PdfReader
    reader=PdfReader(io.BytesIO(payload),strict=False)
    pages=[]
    for page in reader.pages[:1000]:
        pages.append(page.extract_text() or "")
        if sum(map(len,pages))>=MAX_TEXT_CHARS:
            break
    return Extraction(_bounded("\n\n".join(pages)),{"extractor":"pypdf","pages":len(reader.pages)})


def extract_docx(payload:bytes)->Extraction:
    from docx import Document
    _validate_ooxml(payload)
    document=Document(io.BytesIO(payload))
    parts=[p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
            if sum(map(len,parts))>=MAX_TEXT_CHARS:
                break
    return Extraction(_bounded("\n".join(parts)),{"extractor":"python-docx","paragraphs":len(document.paragraphs),"tables":len(document.tables)})


def extract_xlsx(payload:bytes)->Extraction:
    from openpyxl import load_workbook
    _validate_ooxml(payload)
    workbook=load_workbook(io.BytesIO(payload),read_only=True,data_only=True)
    lines=[]
    cells=0
    sheet_names=[]
    for sheet in workbook.worksheets:
        sheet_names.append(sheet.title)
        lines.append(f"[sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            values=["" if value is None else str(value) for value in row]
            if any(values):
                lines.append("\t".join(values))
            cells+=len(row)
            if cells>=MAX_SHEET_CELLS or sum(map(len,lines))>=MAX_TEXT_CHARS:
                break
        if cells>=MAX_SHEET_CELLS or sum(map(len,lines))>=MAX_TEXT_CHARS:
            break
    workbook.close()
    return Extraction(_bounded("\n".join(lines)),{"extractor":"openpyxl","sheets":sheet_names,"cellsRead":cells})


def extract_csv(payload:bytes)->Extraction:
    text=payload.decode("utf-8-sig")
    rows=[]
    for index,row in enumerate(csv.reader(io.StringIO(text))):
        rows.append("\t".join(row))
        if index>=50000 or sum(map(len,rows))>=MAX_TEXT_CHARS:
            break
    return Extraction(_bounded("\n".join(rows)),{"extractor":"csv","rows":len(rows)})


def extract_json(payload:bytes)->Extraction:
    value=json.loads(payload.decode("utf-8"))
    text=json.dumps(value,ensure_ascii=False,indent=2)
    return Extraction(_bounded(text),{"extractor":"json"})


_whisper_model=None

def extract_media(payload:bytes,mime_type:str)->Extraction:
    global _whisper_model
    suffix={
        "audio/wav":".wav","audio/mpeg":".mp3","audio/mp4":".m4a","audio/webm":".webm",
        "video/mp4":".mp4","video/webm":".webm",
    }.get(mime_type,".bin")
    with tempfile.NamedTemporaryFile(suffix=suffix) as src, tempfile.NamedTemporaryFile(suffix=".wav") as wav:
        src.write(payload);src.flush()
        proc=subprocess.run([
            "ffmpeg","-hide_banner","-loglevel","error","-y","-i",src.name,
            "-vn","-ac","1","-ar","16000","-c:a","pcm_s16le",wav.name,
        ],capture_output=True,timeout=180)
        if proc.returncode:
            raise ExtractionError("ARTIFACT_FFMPEG_AUDIO_EXTRACTION_FAILED")
        if _whisper_model is None:
            from faster_whisper import WhisperModel
            _whisper_model=WhisperModel(
                os.getenv("JHADINA_WHISPER_MODEL","small"),
                device=os.getenv("JHADINA_WHISPER_DEVICE","auto"),
                compute_type=os.getenv("JHADINA_WHISPER_COMPUTE","int8"),
            )
        segments,info=_whisper_model.transcribe(wav.name,vad_filter=True,word_timestamps=True)
        text=[]
        segment_rows=[]
        for segment in segments:
            value=segment.text.strip()
            if value:text.append(value)
            segment_rows.append({"startMs":round(segment.start*1000),"endMs":round(segment.end*1000),"text":value})
            if sum(map(len,text))>=MAX_TEXT_CHARS:
                break
        return Extraction(_bounded(" ".join(text)),{
            "extractor":"faster-whisper",
            "language":getattr(info,"language","und"),
            "segments":segment_rows,
        })


def extract(payload:bytes,mime_type:str)->Extraction:
    if mime_type=="application/pdf":return extract_pdf(payload)
    if mime_type==DOCX:return extract_docx(payload)
    if mime_type==XLSX:return extract_xlsx(payload)
    if mime_type=="text/csv":return extract_csv(payload)
    if mime_type=="application/json":return extract_json(payload)
    if mime_type=="text/plain":return Extraction(_bounded(payload.decode("utf-8")),{"extractor":"utf-8"})
    if mime_type in MEDIA_MIME:return extract_media(payload,mime_type)
    raise ExtractionError("ARTIFACT_EXTRACTION_MIME_UNSUPPORTED")
