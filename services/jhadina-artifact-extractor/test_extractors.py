import io
import json
import unittest

from docx import Document
from openpyxl import Workbook

from extractors import ExtractionError,extract


class ExtractorsTest(unittest.TestCase):
    def test_json_and_csv(self):
        self.assertIn('"ok": true',extract(json.dumps({"ok":True}).encode(),"application/json").text)
        self.assertEqual(extract(b"a,b\n1,2\n","text/csv").text,"a\tb\n1\t2")

    def test_docx(self):
        stream=io.BytesIO()
        doc=Document();doc.add_paragraph("Jhadina document body");doc.save(stream)
        result=extract(stream.getvalue(),"application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        self.assertIn("Jhadina document body",result.text)

    def test_xlsx(self):
        stream=io.BytesIO()
        wb=Workbook();ws=wb.active;ws.title="Data";ws.append(["name","value"]);ws.append(["Jhadina",42]);wb.save(stream)
        result=extract(stream.getvalue(),"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("[sheet: Data]",result.text)
        self.assertIn("Jhadina\t42",result.text)

    def test_unsupported_fails_closed(self):
        with self.assertRaisesRegex(ExtractionError,"MIME_UNSUPPORTED"):
            extract(b"x","application/octet-stream")


if __name__=="__main__":
    unittest.main()
